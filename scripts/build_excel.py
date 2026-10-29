"""Construit le classeur Excel décisionnel « Tableau_de_bord_Mali.xlsx ».

Le classeur reprend la base enrichie (géocodage région/cercle, mode opératoire,
fiabilité), calcule toutes les synthèses par formules (elles se recalculent si
la feuille INCIDENTS est modifiée), intègre les cartes de situation et ouvre
la cartographie depuis Excel par liens cliquables (tableau de bord en ligne,
fichier HTML local, OpenStreetMap, Google Maps).

Usage : python scripts/build_excel.py [chemin.xlsx]
Prérequis : python scripts/build_dashboard.py (produit dist/ et dist/cartes/).
"""
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_dashboard import DEFAULT_XLSX, ROOT, build  # noqa: E402

DASHBOARD_URL = "https://claude.ai/artifact/LG5EyCJxmMRJWAq2DKEFbX"
HTML_FILE = "tableau_de_bord_mali.html"
OUT = ROOT / "dist" / "Tableau_de_bord_Mali.xlsx"
MAXR = 500  # plage des formules : lignes 2 à 500 de INCIDENTS (place pour de nouveaux incidents)

FONT = "Arial"
NAVY, TEAL, INK, MUTED, LINE, PALE, PALE2 = "13202C", "0F7C86", "13202C", "566577", "D3DBE3", "EEF2F5", "F7F9FB"
ACTORS = [("JNIM", "DC4A26"), ("EI-Sahel", "8A4FD6"), ("FLA", "D69A12"), ("FAMa / Africa Corps", "2F6FB5"), ("Indéterminé / accident", "7D8A97")]
SEV = {"Critique": "C8102E", "Élevé": "E2701C", "Modéré": "D1A20F", "Faible": "4F8A3C"}
REL = {"Corroboré": "4F8A3C", "Partiellement corroboré": "D1A20F", "Source partisane seule": "E2701C", "Source unique": "C8102E"}

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)


def f(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LINK = f(10, color="0563C1", underline="single")


def banner(ws, title, subtitle, width=12):
    ws.sheet_view.showGridLines = False
    for c in range(1, width + 1):
        for r in (1, 2):
            ws.cell(r, c).fill = fill(NAVY)
    ws["A1"].value = title
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"].value = subtitle
    ws["A2"].font = Font(name=FONT, size=9, color="C9D3DC")
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    back = ws.cell(1, width)
    if ws.title != "ACCUEIL":
        back.value = "◄ Accueil"
        back.hyperlink = "#ACCUEIL!A1"
        back.font = Font(name=FONT, size=9, color="FFFFFF", underline="single")
        back.alignment = Alignment(horizontal="right", vertical="center")


def header_row(ws, row, labels, col=1, color=NAVY):
    for k, label in enumerate(labels):
        c = ws.cell(row, col + k, label)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = fill(color)
        c.alignment = CENTER
        c.border = BOX
    ws.row_dimensions[row].height = 30


def section(ws, row, text, col=1, span=8):
    c = ws.cell(row, col, text)
    c.font = f(12, True, TEAL)
    for k in range(span):
        ws.cell(row, col + k).border = Border(bottom=Side(style="medium", color=TEAL))
    ws.row_dimensions[row].height = 22


def body(c, bold=False, fmt=None, align=None, color=INK, size=10):
    c.font = f(size, bold, color)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    c.alignment = align or Alignment(vertical="top", wrap_text=True)
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def labels(val=False, pct=False):
    return DataLabelList(showVal=val, showPercent=pct, showSerName=False, showCatName=False, showLegendKey=False, showLeaderLines=False)


def style_chart(ch, title, ytitle=None, xtitle=None, w=18, h=8.5):
    ch.title = title
    ch.width, ch.height = w, h
    ch.style = 2
    if ytitle is not None and hasattr(ch, "y_axis"):
        ch.y_axis.title = ytitle
    if xtitle is not None and hasattr(ch, "x_axis"):
        ch.x_axis.title = xtitle
    if hasattr(ch, "y_axis"):
        ch.y_axis.majorGridlines = None
        ch.y_axis.delete = False
        ch.x_axis.delete = False
    ch.legend.position = "b"
    return ch


# --------------------------------------------------------------------------- INCIDENTS
INC_COLS = [
    ("ID", "id", 13), ("Date", "date", 11), ("Semaine ISO", None, 9), ("Période", None, 8),
    ("Pays (géocodé)", "country", 12), ("Région (géocodée)", "region", 15), ("Cercle (géocodé)", "cercle", 14),
    ("Région déclarée", "region_reported", 18), ("Localité", "place", 24), ("Latitude", "lat", 9), ("Longitude", "lon", 9),
    ("Précision géo.", "precision", 12), ("Typologie (source)", "typology", 24), ("Mode opératoire", "mode", 20),
    ("Instigateur", "actor", 16), ("Instigateur (détail)", "actor_raw", 22), ("Cible", "target", 16), ("Cible (détail)", "target_raw", 22),
    ("Description", "desc", 60), ("Morts civils", "civ_k", 8), ("Blessés civils", "civ_w", 8), ("Morts FDS (FAMa)", "fama_k", 8),
    ("Morts Africa Corps", "ac_k", 8), ("Morts groupes armés", "ga_k", 8), ("Total tués", None, 8), ("Capturés", "pow", 8),
    ("Véh./aéronefs détruits", "veh_d", 9), ("Véh./aéronefs saisis", "veh_s", 9), ("Armement / équipement saisi", "equipment", 26),
    ("Impact", "impact", 26), ("Distance Bamako (km)", None, 10), ("Dans le rayon Bamako", None, 9), ("Fiabilité", "reliability", 16),
    ("Bilan revendiqué seul", None, 10), ("Sources", "sources", 30), ("Carte OSM", None, 11), ("Google Maps", None, 11), ("Tableau de bord", None, 12),
]
COL = {name: get_column_letter(k + 1) for k, (name, _, _) in enumerate(INC_COLS)}


def sheet_incidents(wb, incidents):
    ws = wb.create_sheet("INCIDENTS")
    ws.sheet_properties.tabColor = TEAL
    header_row(ws, 1, [c[0] for c in INC_COLS])
    for k, (_, _, w) in enumerate(INC_COLS):
        ws.column_dimensions[get_column_letter(k + 1)].width = w
    L = COL
    for r, i in enumerate(incidents, start=2):
        for k, (name, key, _) in enumerate(INC_COLS):
            c = ws.cell(r, k + 1)
            if key == "date":
                c.value = date.fromisoformat(i["date"])
                c.number_format = "DD/MM/YYYY"
            elif key == "sources":
                c.value = ", ".join(s["name"] for s in i["sources"])
            elif key:
                c.value = i[key]
            c.font = f(9)
            c.alignment = Alignment(vertical="top", wrap_text=key in ("desc", "equipment", "impact", "sources", "typology", "actor_raw", "target_raw", "place", "region_reported"))
            c.border = Border(bottom=thin)
        lat, lon = f"{L['Latitude']}{r}", f"{L['Longitude']}{r}"
        ws[f"{L['Semaine ISO']}{r}"] = f"=WEEKNUM({L['Date']}{r},21)"
        ws[f"{L['Période']}{r}"] = f'=IF({L["Date"]}{r}<P_DEBUT_P2,"P1","P2")'
        ws[f"{L['Total tués']}{r}"] = f"={L['Morts civils']}{r}+{L['Morts FDS (FAMa)']}{r}+{L['Morts Africa Corps']}{r}+{L['Morts groupes armés']}{r}"
        ws[f"{L['Distance Bamako (km)']}{r}"] = (f"=ROUND(2*6371*ASIN(SQRT(SIN(RADIANS({lat}-LAT_BAMAKO)/2)^2+COS(RADIANS(LAT_BAMAKO))*COS(RADIANS({lat}))"
                                              f"*SIN(RADIANS({lon}-LON_BAMAKO)/2)^2)),0)")
        ws[f"{L['Dans le rayon Bamako']}{r}"] = f'=IF({L["Distance Bamako (km)"]}{r}<=SEUIL_KM,"Oui","Non")'
        ws[f"{L['Bilan revendiqué seul']}{r}"] = "Oui" if i["claim_only"] else "Non"
        # Coordonnées converties en texte à point décimal : le séparateur « , » d'un Excel français casserait l'URL.
        tl, tn = f'SUBSTITUTE({lat}&"",",",".")', f'SUBSTITUTE({lon}&"",",",".")'
        ws[f"{L['Carte OSM']}{r}"] = f'=HYPERLINK("https://www.openstreetmap.org/?mlat="&{tl}&"&mlon="&{tn}&"#map=11/"&{tl}&"/"&{tn},"OSM ↗")'
        ws[f"{L['Google Maps']}{r}"] = f'=HYPERLINK("https://www.google.com/maps/search/?api=1&query="&{tl}&"%2C"&{tn},"Google ↗")'
        ws[f"{L['Tableau de bord']}{r}"] = f'=HYPERLINK(URL_TABLEAU&"#"&{L["ID"]}{r},"Dossier ↗")'
        for name in ("Carte OSM", "Google Maps", "Tableau de bord"):
            ws[f"{L[name]}{r}"].font = f(9, color="0563C1", underline="single")
        for name in ("Semaine ISO", "Période", "Total tués", "Distance Bamako (km)", "Dans le rayon Bamako", "Bilan revendiqué seul"):
            c = ws[f"{L[name]}{r}"]
            c.font = f(9, name == "Total tués")
            c.alignment = Alignment(horizontal="center", vertical="top")
            c.border = Border(bottom=thin)
        for name in ("Latitude", "Longitude"):
            ws[f"{L[name]}{r}"].number_format = "0.0000"
        ws.row_dimensions[r].height = 48
    last = len(incidents) + 1
    tab = Table(displayName="tblIncidents", ref=f"A1:{get_column_letter(len(INC_COLS))}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "B2"
    rng = lambda c: f"{L[c]}2:{L[c]}{MAXR}"  # noqa: E731
    ws.conditional_formatting.add(rng("Total tués"), DataBarRule(start_type="num", start_value=0, end_type="max", color="E2701C", showValue=True))
    ws.conditional_formatting.add(rng("Morts civils"), CellIsRule(operator="greaterThan", formula=["0"], fill=fill("F9D3D8"), font=Font(name=FONT, bold=True, color="C8102E")))
    ws.conditional_formatting.add(rng("Blessés civils"), CellIsRule(operator="greaterThan", formula=["0"], fill=fill("F9D3D8"), font=Font(name=FONT, color="C8102E")))
    ws.conditional_formatting.add(rng("Dans le rayon Bamako"), CellIsRule(operator="equal", formula=['"Oui"'], fill=fill("FBE3D0"), font=Font(name=FONT, bold=True, color="B4501A")))
    ws.conditional_formatting.add(rng("Bilan revendiqué seul"), CellIsRule(operator="equal", formula=['"Oui"'], font=Font(name=FONT, bold=True, color="E2701C")))
    for k, (rel, col) in enumerate(REL.items()):
        ws.conditional_formatting.add(rng("Fiabilité"), CellIsRule(operator="equal", formula=[f'"{rel}"'], font=Font(name=FONT, bold=True, color=col)))
    for a, col in ACTORS:
        ws.conditional_formatting.add(rng("Instigateur"), CellIsRule(operator="equal", formula=[f'"{a}"'], font=Font(name=FONT, bold=True, color=col)))
    ws["A1"].comment = Comment("Base source : DATA_CARTOGRAPHIE (compilations Naïm TOURÉ). Colonnes « géocodées » : jointure spatiale des coordonnées sur geoBoundaries. "
                               "Les colonnes en formule (semaine, période, total, distance, liens) se recalculent automatiquement.", "Observatoire")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    return last


def inc_names(wb):
    names = {"ID": "inc_id", "Date": "inc_date", "Période": "inc_periode", "Région (géocodée)": "inc_region", "Cercle (géocodé)": "inc_cercle",
             "Mode opératoire": "inc_mode", "Instigateur": "inc_acteur", "Cible": "inc_cible", "Morts civils": "inc_civ_k", "Blessés civils": "inc_civ_w",
             "Morts FDS (FAMa)": "inc_fama", "Morts Africa Corps": "inc_ac", "Morts groupes armés": "inc_ga", "Total tués": "inc_tues", "Capturés": "inc_pow",
             "Véh./aéronefs détruits": "inc_vehd", "Véh./aéronefs saisis": "inc_vehs", "Distance Bamako (km)": "inc_dist", "Dans le rayon Bamako": "inc_rayon",
             "Fiabilité": "inc_fiab", "Sources": "inc_sources", "Bilan revendiqué seul": "inc_revend", "Précision géo.": "inc_prec"}
    for col, n in names.items():
        add_name(wb, n, f"INCIDENTS!${COL[col]}$2:${COL[col]}${MAXR}")


# --------------------------------------------------------------------------- PARAMETRES
def sheet_params(wb, meta):
    ws = wb.create_sheet("PARAMETRES")
    banner(ws, "Paramètres du classeur", "Cellules en bleu : valeurs modifiables. Toutes les synthèses s'y réfèrent.", 6)
    widths(ws, {"A": 3, "B": 42, "C": 30, "D": 60, "E": 3, "F": 12})
    rows = [
        ("Période d'analyse", None, None, None),
        ("Début de la période", date.fromisoformat(meta["start"]), "P_DEBUT", "Premier incident de la base."),
        ("Fin de la période", date.fromisoformat(meta["end"]), "P_FIN", "Dernier incident de la base. Sert aussi de référence pour l'activité récente."),
        ("Début de la période P2", date.fromisoformat(meta["split"]), "P_DEBUT_P2", "P1 = avant cette date, P2 = à partir de cette date (deux périodes de durée égale)."),
        ("Pression sur Bamako", None, None, None),
        ("Latitude de Bamako", meta["bamako"]["lat"], "LAT_BAMAKO", "Centre-ville, WGS84."),
        ("Longitude de Bamako", meta["bamako"]["lon"], "LON_BAMAKO", "Centre-ville, WGS84."),
        ("Rayon de pression (km)", 300, "SEUIL_KM", "Incidents comptés « dans le rayon de Bamako »."),
        ("Accès à la cartographie", None, None, None),
        ("Tableau de bord en ligne (URL)", DASHBOARD_URL, "URL_TABLEAU", "Lien privé : partagez-le depuis le menu « Partager » de la page."),
        ("Fichier HTML local", HTML_FILE, "FICHIER_HTML", "Doit se trouver dans le même dossier que ce classeur."),
        ("Indice de menace régional", None, None, None),
        ("Fenêtre d'activité récente (jours)", 21, "FEN_RECENTE", "Incidents des N derniers jours de la période."),
        ("Poids : fréquence des incidents", 0.35, "W_FREQ", "La somme des quatre poids doit faire 100 %."),
        ("Poids : létalité (échelle logarithmique)", 0.35, "W_LETAL", ""),
        ("Poids : victimes civiles (tués + blessés)", 0.15, "W_CIVILS", ""),
        ("Poids : activité récente", 0.15, "W_RECENT", ""),
        ("Seuil niveau « Critique »", 70, "S_CRITIQUE", "Indice ≥ seuil."),
        ("Seuil niveau « Élevé »", 45, "S_ELEVE", ""),
        ("Seuil niveau « Modéré »", 25, "S_MODERE", "En dessous : « Faible »."),
    ]
    r = 4
    for label, val, name, note in rows:
        if val is None:
            section(ws, r, label, 2, 3)
            r += 1
            continue
        body(ws.cell(r, 2, label))
        c = body(ws.cell(r, 3, val), color="0000FF", bold=True)
        if isinstance(val, date):
            c.number_format = "DD/MM/YYYY"
        elif isinstance(val, float) and 0 < val < 1:
            c.number_format = "0%"
        body(ws.cell(r, 4, note), color=MUTED, size=9)
        add_name(wb, name, f"PARAMETRES!$C${r}")
        r += 1
    body(ws.cell(r, 2, "Contrôle : somme des poids"), bold=True)
    c = body(ws.cell(r, 3, "=W_FREQ+W_LETAL+W_CIVILS+W_RECENT"), bold=True, fmt="0%")
    ws.conditional_formatting.add(f"C{r}", CellIsRule(operator="notEqual", formula=["1"], fill=fill("F9D3D8")))
    body(ws.cell(r, 4, "Doit afficher 100 %. Fond rouge sinon."), color=MUTED, size=9)


# --------------------------------------------------------------------------- REGIONS
def sheet_regions(wb, incidents):
    ws = wb.create_sheet("REGIONS")
    banner(ws, "Synthèse régionale et indice de menace", "Toutes les valeurs sont calculées par formule à partir de la feuille INCIDENTS.", 16)
    regs = sorted({i["region"] for i in incidents}, key=lambda r: (-sum(i["region"] == r for i in incidents), r))
    hdr = ["Région", "Incidents", "Part", "Tués", "FDS + Africa Corps tués", "Civils tués + blessés", "Incidents récents", "Indice (0-100)", "Niveau", "Acteur dominant"] + [a for a, _ in ACTORS]
    header_row(ws, 4, hdr, 2)
    for k, (a, col) in enumerate(ACTORS):
        ws.cell(4, 12 + k).fill = fill(col)
    n0, n1 = 5, 4 + len(regs)
    for r, reg in enumerate(regs, start=n0):
        body(ws.cell(r, 2, reg), bold=True)
        body(ws.cell(r, 3, f"=COUNTIFS(inc_region,$B{r})"), fmt="0", align=CENTER)
        body(ws.cell(r, 4, f"=IF(SUM($C${n0}:$C${n1})=0,0,C{r}/SUM($C${n0}:$C${n1}))"), fmt="0%", align=CENTER)
        body(ws.cell(r, 5, f"=SUMIFS(inc_tues,inc_region,$B{r})"), fmt="0", align=CENTER)
        body(ws.cell(r, 6, f"=SUMIFS(inc_fama,inc_region,$B{r})+SUMIFS(inc_ac,inc_region,$B{r})"), fmt="0", align=CENTER)
        body(ws.cell(r, 7, f"=SUMIFS(inc_civ_k,inc_region,$B{r})+SUMIFS(inc_civ_w,inc_region,$B{r})"), fmt="0", align=CENTER)
        body(ws.cell(r, 8, f'=COUNTIFS(inc_region,$B{r},inc_date,">="&(P_FIN-FEN_RECENTE+1))'), fmt="0", align=CENTER)
        body(ws.cell(r, 9, (f"=IF(C{r}=0,0,ROUND(100*(W_FREQ*C{r}/MAX($C${n0}:$C${n1})+W_LETAL*LN(1+E{r})/LN(1+MAX(1,MAX($E${n0}:$E${n1})))"
                            f"+W_CIVILS*IF(MAX($G${n0}:$G${n1})=0,0,G{r}/MAX($G${n0}:$G${n1}))+W_RECENT*H{r}/C{r}),0))")), bold=True, fmt="0", align=CENTER)
        body(ws.cell(r, 10, f'=IF(I{r}>=S_CRITIQUE,"Critique",IF(I{r}>=S_ELEVE,"Élevé",IF(I{r}>=S_MODERE,"Modéré","Faible")))'), bold=True, align=CENTER)
        body(ws.cell(r, 11, f"=IF(C{r}=0,\"—\",INDEX($L$4:$P$4,MATCH(MAX(L{r}:P{r}),L{r}:P{r},0)))"), align=CENTER)
        for k, (a, _) in enumerate(ACTORS):
            body(ws.cell(r, 12 + k, f"=COUNTIFS(inc_region,$B{r},inc_acteur,{get_column_letter(12 + k)}$4)"), fmt="0;-0;\"·\"", align=CENTER)
    t = n1 + 1
    body(ws.cell(t, 2, "Total"), bold=True).fill = fill(PALE)
    for c in range(3, 17):
        if c in (9, 10, 11):
            ws.cell(t, c).fill = fill(PALE)
            ws.cell(t, c).border = BOX
            continue
        L = get_column_letter(c)
        cell = body(ws.cell(t, c, f"=SUM({L}{n0}:{L}{n1})"), bold=True, fmt="0%" if c == 4 else "0", align=CENTER)
        cell.fill = fill(PALE)
    ws.conditional_formatting.add(f"I{n0}:I{n1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="C8102E", showValue=True))
    for lvl, col in SEV.items():
        ws.conditional_formatting.add(f"J{n0}:J{n1}", CellIsRule(operator="equal", formula=[f'"{lvl}"'], fill=fill(col), font=Font(name=FONT, bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(f"L{n0}:P{n1}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="F0A07A"))
    widths(ws, {"A": 2, "B": 26, "C": 10, "D": 8, "E": 9, "F": 12, "G": 12, "H": 11, "I": 12, "J": 11, "K": 20, "L": 9, "M": 9, "N": 9, "O": 11, "P": 11})
    ws.freeze_panes = "C5"
    n = t + 2
    ws.cell(n, 2, "Méthode : indice = 100 × (W_FREQ × incidents / max + W_LETAL × ln(1+tués) / ln(1+max) + W_CIVILS × civils / max + W_RECENT × part d'incidents récents). "
                  "Poids et seuils modifiables dans PARAMETRES. Les régions suivent le référentiel geoBoundaries (8 régions + Bamako, antérieur à la réforme de 2023).").font = f(9, color=MUTED, italic=True)
    ws.merge_cells(start_row=n, start_column=2, end_row=n, end_column=16)
    ws.cell(n, 2).alignment = WRAP
    ws.row_dimensions[n].height = 30
    ch = style_chart(BarChart(), "Indice de menace par région", "Indice (0-100)")
    ch.type = "bar"
    ch.add_data(Reference(ws, min_col=9, min_row=4, max_row=n1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    ch.series[0].graphicalProperties.solidFill = "C8102E"
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
    ch.dataLabels = labels(val=True)
    ch.legend = None
    ws.add_chart(ch, f"B{n + 2}")
    ch2 = style_chart(BarChart(), "Incidents par région et par acteur", "Incidents")
    ch2.type, ch2.grouping, ch2.overlap = "bar", "stacked", 100
    for k, (a, col) in enumerate(ACTORS):
        ch2.add_data(Reference(ws, min_col=12 + k, min_row=4, max_row=n1), titles_from_data=True)
        ch2.series[-1].graphicalProperties.solidFill = col
    ch2.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    ch2.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch2, f"I{n + 2}")
    return regs


# --------------------------------------------------------------------------- ACTEURS
def sheet_actors(wb, incidents):
    ws = wb.create_sheet("ACTEURS")
    banner(ws, "Synthèse par acteur instigateur", "« Tués lors de ses actions » inclut les pertes subies par l'acteur lui-même lors des incidents qu'il initie.", 15)
    hdr = ["Acteur", "Incidents", "Part", "Tués lors de ses actions", "FDS (FAMa) tués", "Africa Corps tués", "Civils tués", "Civils blessés",
           "Combattants armés tués", "Capturés", "Véh. détruits", "Véh. saisis", "Incidents P1", "Incidents P2", "Évolution P1 → P2"]
    header_row(ws, 4, hdr, 2)
    keys = ["inc_tues", "inc_fama", "inc_ac", "inc_civ_k", "inc_civ_w", "inc_ga", "inc_pow", "inc_vehd", "inc_vehs"]
    n0 = 5
    n1 = n0 + len(ACTORS) - 1
    for r, (a, col) in enumerate(ACTORS, start=n0):
        c = body(ws.cell(r, 2, a), bold=True, color="FFFFFF")
        c.fill = fill(col)
        body(ws.cell(r, 3, f"=COUNTIFS(inc_acteur,$B{r})"), fmt="0", align=CENTER)
        body(ws.cell(r, 4, f"=IF(SUM($C${n0}:$C${n1})=0,0,C{r}/SUM($C${n0}:$C${n1}))"), fmt="0%", align=CENTER)
        for k, nm in enumerate(keys):
            body(ws.cell(r, 5 + k, f"=SUMIFS({nm},inc_acteur,$B{r})"), fmt="0;-0;\"·\"", align=CENTER)
        body(ws.cell(r, 14, f'=COUNTIFS(inc_acteur,$B{r},inc_periode,"P1")'), fmt="0", align=CENTER)
        body(ws.cell(r, 15, f'=COUNTIFS(inc_acteur,$B{r},inc_periode,"P2")'), fmt="0", align=CENTER)
        body(ws.cell(r, 16, f'=IF(N{r}=0,IF(O{r}=0,"—","nouveau"),(O{r}-N{r})/N{r})'), fmt='+0%;-0%;0%', align=CENTER)
    t = n1 + 1
    body(ws.cell(t, 2, "Total"), bold=True).fill = fill(PALE)
    for c in range(3, 16):
        L = get_column_letter(c)
        body(ws.cell(t, c, f"=SUM({L}{n0}:{L}{n1})"), bold=True, fmt="0%" if c == 4 else "0", align=CENTER).fill = fill(PALE)
    body(ws.cell(t, 16, f'=IF(N{t}=0,"—",(O{t}-N{t})/N{t})'), bold=True, fmt='+0%;-0%;0%', align=CENTER).fill = fill(PALE)
    ws.conditional_formatting.add(f"E{n0}:E{n1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="E2701C", showValue=True))
    ws.conditional_formatting.add(f"H{n0}:I{n1}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("F9D3D8"), font=Font(name=FONT, bold=True, color="C8102E")))
    widths(ws, {"A": 2, "B": 24, **{get_column_letter(c): 11 for c in range(3, 17)}})
    ws.column_dimensions["E"].width = 14
    # Matrice acteur × mode opératoire
    modes = sorted({i["mode"] for i in incidents})
    m0 = t + 3
    section(ws, m0 - 1, "Matrice acteur × mode opératoire (nombre d'incidents)", 2, len(modes) + 2)
    header_row(ws, m0, ["Acteur"] + modes + ["Total"], 2, TEAL)
    for r, (a, col) in enumerate(ACTORS, start=m0 + 1):
        c = body(ws.cell(r, 2, a), bold=True)
        c.font = f(10, True, col)
        for k, m in enumerate(modes):
            L = get_column_letter(3 + k)
            body(ws.cell(r, 3 + k, f"=COUNTIFS(inc_acteur,$B{r},inc_mode,{L}${m0})"), fmt="0;-0;\"·\"", align=CENTER)
        body(ws.cell(r, 3 + len(modes), f"=SUM(C{r}:{get_column_letter(2 + len(modes))}{r})"), bold=True, align=CENTER)
    ws.conditional_formatting.add(f"C{m0 + 1}:{get_column_letter(2 + len(modes))}{m0 + len(ACTORS)}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=2, mid_color="F6C99B", end_type="max", end_color="C8102E"))
    ws.row_dimensions[m0].height = 54
    g = m0 + len(ACTORS) + 3
    pie = PieChart()
    pie.title = "Part des incidents par acteur"
    pie.add_data(Reference(ws, min_col=3, min_row=4, max_row=n1), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    from openpyxl.chart.series import DataPoint
    for k, (_, col) in enumerate(ACTORS):
        pt = DataPoint(idx=k)
        pt.graphicalProperties.solidFill = col
        pie.series[0].dPt.append(pt)
    pie.dataLabels = labels(pct=True)
    pie.width, pie.height = 12, 8.5
    ws.add_chart(pie, f"B{g}")
    ch = style_chart(BarChart(), "Pertes infligées lors des actions de chaque acteur", "Tués")
    ch.grouping, ch.overlap = "stacked", 100
    for k, (lab, colr) in enumerate([("FDS (FAMa) tués", "2F6FB5"), ("Africa Corps tués", "7D8A97"), ("Civils tués", "C8102E"), ("Combattants armés tués", "DC4A26")]):
        idx = hdr.index(lab) + 2
        ch.add_data(Reference(ws, min_col=idx, min_row=4, max_row=n1), titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = colr
    ch.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    ws.add_chart(ch, f"H{g}")


# --------------------------------------------------------------------------- CHRONOLOGIE
def sheet_chrono(wb, meta):
    ws = wb.create_sheet("CHRONOLOGIE")
    banner(ws, "Chronologie et comparaison des périodes", "Semaines ISO (lundi → dimanche). P1 / P2 : deux périodes de durée égale (voir PARAMETRES).", 15)
    d0, d1 = date.fromisoformat(meta["start"]), date.fromisoformat(meta["end"])
    monday = d0 - timedelta(days=d0.weekday())
    weeks = []
    while monday <= d1:
        weeks.append(monday)
        monday += timedelta(days=7)
    hdr = ["Semaine du", "Au", "N° ISO"] + [a for a, _ in ACTORS] + ["Total incidents", "Tués", "FDS + AC tués", "Civils tués", "Cumul tués"]
    header_row(ws, 4, hdr, 2)
    for k, (a, col) in enumerate(ACTORS):
        ws.cell(4, 5 + k).fill = fill(col)
    n0 = 5
    n1 = n0 + len(weeks) - 1
    for r, w in enumerate(weeks, start=n0):
        body(ws.cell(r, 2, w), fmt="DD/MM/YYYY", align=CENTER)
        body(ws.cell(r, 3, f"=B{r}+6"), fmt="DD/MM/YYYY", align=CENTER)
        body(ws.cell(r, 4, f'="S"&WEEKNUM(B{r},21)'), align=CENTER)
        for k in range(len(ACTORS)):
            L = get_column_letter(5 + k)
            body(ws.cell(r, 5 + k, f'=COUNTIFS(inc_acteur,{L}$4,inc_date,">="&$B{r},inc_date,"<="&$C{r})'), fmt="0;-0;\"·\"", align=CENTER)
        body(ws.cell(r, 10, f"=SUM(E{r}:I{r})"), bold=True, align=CENTER)
        body(ws.cell(r, 11, f'=SUMIFS(inc_tues,inc_date,">="&$B{r},inc_date,"<="&$C{r})'), align=CENTER)
        body(ws.cell(r, 12, f'=SUMIFS(inc_fama,inc_date,">="&$B{r},inc_date,"<="&$C{r})+SUMIFS(inc_ac,inc_date,">="&$B{r},inc_date,"<="&$C{r})'), align=CENTER)
        body(ws.cell(r, 13, f'=SUMIFS(inc_civ_k,inc_date,">="&$B{r},inc_date,"<="&$C{r})'), align=CENTER)
        body(ws.cell(r, 14, f"=SUM($K${n0}:K{r})"), align=CENTER)
    ws.conditional_formatting.add(f"J{n0}:J{n1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="0F7C86", showValue=True))
    ws.conditional_formatting.add(f"M{n0}:M{n1}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill("F9D3D8"), font=Font(name=FONT, bold=True, color="C8102E")))
    widths(ws, {"A": 2, "B": 12, "C": 12, "D": 8, **{get_column_letter(c): 11 for c in range(5, 15)}, "P": 2, "Q": 30, "R": 11, "S": 11, "T": 11, "U": 11})
    # Comparaison P1 / P2
    section(ws, 4, "Comparaison P1 / P2", 17, 5)
    header_row(ws, 5, ["Indicateur", "P1", "P2", "Écart", "Évolution"], 17, TEAL)
    comp = [("Incidents", None), ("Tués (total)", "inc_tues"), ("FDS (FAMa) tués", "inc_fama"), ("Africa Corps tués", "inc_ac"), ("Combattants armés tués", "inc_ga"),
            ("Civils tués", "inc_civ_k"), ("Civils blessés", "inc_civ_w"), ("Capturés", "inc_pow"), ("Véhicules / aéronefs détruits", "inc_vehd"), ("Incidents dans le rayon de Bamako", "rayon")]
    for r, (lab, nm) in enumerate(comp, start=6):
        body(ws.cell(r, 17, lab), bold=True)
        for c, p in ((18, "P1"), (19, "P2")):
            if nm is None:
                fml = f'=COUNTIFS(inc_periode,"{p}")'
            elif nm == "rayon":
                fml = f'=COUNTIFS(inc_periode,"{p}",inc_rayon,"Oui")'
            else:
                fml = f'=SUMIFS({nm},inc_periode,"{p}")'
            body(ws.cell(r, c, fml), fmt="0", align=CENTER)
        body(ws.cell(r, 20, f"=S{r}-R{r}"), fmt="+0;-0;0", align=CENTER)
        body(ws.cell(r, 21, f'=IF(R{r}=0,"—",(S{r}-R{r})/R{r})'), fmt="+0%;-0%;0%", align=CENTER, bold=True)
    ws.conditional_formatting.add("U6:U15", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, bold=True, color="C8102E")))
    ws.conditional_formatting.add("U6:U15", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="4F8A3C")))
    ws.cell(16, 17, "Rouge : hausse ; vert : baisse. La baisse des pertes FAMa en P2 tient surtout à l'effet Dioura (150 tués le 10/09).").font = f(8, color=MUTED, italic=True)
    g = max(n1, 16) + 2
    bar = style_chart(BarChart(), "Rythme opérationnel hebdomadaire par acteur", "Incidents", w=22)
    bar.grouping, bar.overlap = "stacked", 100
    for k, (a, col) in enumerate(ACTORS):
        bar.add_data(Reference(ws, min_col=5 + k, min_row=4, max_row=n1), titles_from_data=True)
        bar.series[-1].graphicalProperties.solidFill = col
    bar.set_categories(Reference(ws, min_col=4, min_row=n0, max_row=n1))
    line = LineChart()
    line.add_data(Reference(ws, min_col=11, min_row=4, max_row=n1), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.title = "Tués"
    line.y_axis.crosses = "max"
    line.series[0].graphicalProperties.line.solidFill = "13202C"
    line.series[0].graphicalProperties.line.width = 28000
    line.series[0].smooth = True
    bar += line
    ws.add_chart(bar, f"B{g}")
    ch = style_chart(BarChart(), "Bilan humain hebdomadaire", "Tués", w=16)
    ch.grouping, ch.overlap = "stacked", 100
    ch.add_data(Reference(ws, min_col=12, min_row=4, max_row=n1), titles_from_data=True)
    ch.series[-1].graphicalProperties.solidFill = "2F6FB5"
    ch.add_data(Reference(ws, min_col=13, min_row=4, max_row=n1), titles_from_data=True)
    ch.series[-1].graphicalProperties.solidFill = "C8102E"
    ch.set_categories(Reference(ws, min_col=4, min_row=n0, max_row=n1))
    ws.add_chart(ch, f"L{g}")


# --------------------------------------------------------------------------- MODES & CIBLES
def sheet_modes(wb, incidents):
    ws = wb.create_sheet("MODES_CIBLES")
    banner(ws, "Modes opératoires et cibles", "Fréquence, létalité et matrice mode × cible.", 12)
    modes = sorted({i["mode"] for i in incidents}, key=lambda m: -sum(i["mode"] == m for i in incidents))
    targets = sorted({i["target"] for i in incidents}, key=lambda t: -sum(i["target"] == t for i in incidents))
    header_row(ws, 4, ["Mode opératoire", "Incidents", "Part", "Tués", "Tués par incident", "Civils tués", "Dans le rayon Bamako"], 2)
    n0, n1 = 5, 4 + len(modes)
    for r, m in enumerate(modes, start=n0):
        body(ws.cell(r, 2, m), bold=True)
        body(ws.cell(r, 3, f"=COUNTIFS(inc_mode,$B{r})"), align=CENTER)
        body(ws.cell(r, 4, f"=IF(SUM($C${n0}:$C${n1})=0,0,C{r}/SUM($C${n0}:$C${n1}))"), fmt="0%", align=CENTER)
        body(ws.cell(r, 5, f"=SUMIFS(inc_tues,inc_mode,$B{r})"), align=CENTER)
        body(ws.cell(r, 6, f"=IF(C{r}=0,0,E{r}/C{r})"), fmt="0.0", align=CENTER)
        body(ws.cell(r, 7, f"=SUMIFS(inc_civ_k,inc_mode,$B{r})"), fmt="0;-0;\"·\"", align=CENTER)
        body(ws.cell(r, 8, f'=COUNTIFS(inc_mode,$B{r},inc_rayon,"Oui")'), fmt="0;-0;\"·\"", align=CENTER)
    ws.conditional_formatting.add(f"C{n0}:C{n1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="0F7C86", showValue=True))
    ws.conditional_formatting.add(f"F{n0}:F{n1}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="E2701C"))
    widths(ws, {"A": 2, "B": 28, "C": 11, "D": 9, "E": 9, "F": 12, "G": 11, "H": 12, "I": 3, "J": 28, **{get_column_letter(c): 13 for c in range(11, 11 + len(targets) + 1)}})
    header_row(ws, 4, ["Mode × cible"] + targets, 10, TEAL)
    for r, m in enumerate(modes, start=n0):
        body(ws.cell(r, 10, f"=B{r}"), bold=True)
        for k, t in enumerate(targets):
            L = get_column_letter(11 + k)
            body(ws.cell(r, 11 + k, f"=COUNTIFS(inc_mode,$J{r},inc_cible,{L}$4)"), fmt="0;-0;\"·\"", align=CENTER)
    ws.conditional_formatting.add(f"K{n0}:{get_column_letter(10 + len(targets))}{n1}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", mid_type="num", mid_value=2, mid_color="F6C99B", end_type="max", end_color="C8102E"))
    g = n1 + 3
    ch = style_chart(BarChart(), "Fréquence des modes opératoires", "Incidents", w=16, h=9)
    ch.type = "bar"
    ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=n1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    ch.series[0].graphicalProperties.solidFill = TEAL
    ch.x_axis.scaling.orientation = "maxMin"
    ch.legend = None
    ws.add_chart(ch, f"B{g}")
    ch2 = style_chart(BarChart(), "Létalité moyenne par mode (tués / incident)", "Tués par incident", w=16, h=9)
    ch2.type = "bar"
    ch2.add_data(Reference(ws, min_col=6, min_row=4, max_row=n1), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=n0, max_row=n1))
    ch2.series[0].graphicalProperties.solidFill = "C8102E"
    ch2.x_axis.scaling.orientation = "maxMin"
    ch2.legend = None
    ws.add_chart(ch2, f"J{g}")


# --------------------------------------------------------------------------- BAMAKO
def sheet_bamako(wb, incidents):
    ws = wb.create_sheet("PRESSION_BAMAKO")
    banner(ws, "Pression sur Bamako", "Distance orthodromique de chaque incident au centre de Bamako (formule de haversine, voir INCIDENTS).", 10)
    bands = [(0, 50), (50, 150), (150, 300), (300, 500), (500, 800), (800, 2000)]
    header_row(ws, 4, ["Tranche de distance", "De (km)", "À (km)", "Incidents", "dont JNIM", "Tués", "Part des incidents"], 2)
    for r, (lo, hi) in enumerate(bands, start=5):
        body(ws.cell(r, 2, f'=C{r}&" – "&D{r}&" km"' if hi < 2000 else f'="Plus de "&C{r}&" km"'), bold=True)
        body(ws.cell(r, 3, lo), color="0000FF", align=CENTER)
        body(ws.cell(r, 4, hi), color="0000FF", align=CENTER)
        body(ws.cell(r, 5, f'=COUNTIFS(inc_dist,">="&C{r},inc_dist,"<"&D{r})'), bold=True, align=CENTER)
        body(ws.cell(r, 6, f'=COUNTIFS(inc_dist,">="&C{r},inc_dist,"<"&D{r},inc_acteur,"JNIM")'), align=CENTER)
        body(ws.cell(r, 7, f'=SUMIFS(inc_tues,inc_dist,">="&C{r},inc_dist,"<"&D{r})'), align=CENTER)
        body(ws.cell(r, 8, f"=IF(SUM($E$5:$E$10)=0,0,E{r}/SUM($E$5:$E$10))"), fmt="0%", align=CENTER)
    ws.conditional_formatting.add("E5:E10", DataBarRule(start_type="num", start_value=0, end_type="max", color="C8102E", showValue=True))
    widths(ws, {"A": 2, "B": 24, "C": 10, "D": 10, "E": 11, "F": 11, "G": 10, "H": 12, "I": 3, "J": 14, "K": 12, "L": 30, "M": 18, "N": 20, "O": 10, "P": 10})
    ws.cell(11, 2, "Bornes « De / À » modifiables (en bleu).").font = f(8, color=MUTED, italic=True)
    section(ws, 13, "Incidents dans le rayon de Bamako (à la génération du classeur, triés par distance)", 2, 8)
    header_row(ws, 14, ["ID", "Date", "Localité", "Cercle", "Instigateur", "Mode", "Distance (km)", "Tués"], 2, TEAL)
    near = sorted((i for i in incidents if i["dist_bamako"] <= 300), key=lambda i: i["dist_bamako"])
    for r, i in enumerate(near, start=15):
        src = 2 + incidents.index(i)
        for k, col in enumerate(["ID", "Date", "Localité", "Cercle (géocodé)", "Instigateur", "Mode opératoire", "Distance Bamako (km)", "Total tués"]):
            c = body(ws.cell(r, 2 + k, f"=INCIDENTS!{COL[col]}{src}"), align=CENTER if k in (1, 6, 7) else None)
            if col == "Date":
                c.number_format = "DD/MM/YYYY"
        ws.cell(r, 2).hyperlink = f"#INCIDENTS!A{src}"
        ws.cell(r, 2).font = LINK
    ch = style_chart(BarChart(), "Incidents par tranche de distance à Bamako", "Incidents", w=15, h=8)
    ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=10), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=6, min_row=4, max_row=10), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=10))
    ch.series[0].graphicalProperties.solidFill = "13202C"
    ch.series[1].graphicalProperties.solidFill = "DC4A26"
    ws.add_chart(ch, "J4")


# --------------------------------------------------------------------------- EVALUATION
JUDGEMENTS = [
    ("Campagne d'usure méthodique du JNIM", "Le JNIM vise les emprises isolées et les axes du Mali central. La prise du camp de Dioura (10/09 : 150 FAMa tués, 93 capturés, 115 véhicules détruits) concentre à elle seule 72 % des pertes FAMa de la période.", "Modérée (bilan issu en partie d'AzZalaqa)"),
    ("Étau autour de Bamako", "Kassela (32 km), axe Mourdja–Kaloumba, Ségou–Fana, Dougabougou–Markala : l'incendie de citernes et les embuscades de convois relèvent très probablement d'une stratégie de blocus économique de la capitale.", "Élevée"),
    ("Ouverture d'un front sud", "La prise d'un poste à Kourouma (Sikasso, 07/10, ≈ 46 km de la frontière burkinabè) indique une capacité d'action dans le couloir Sikasso–Bobo-Dioulasso, à la porte des États côtiers.", "Modérée (un seul événement)"),
    ("Rivalité JNIM – EI-Sahel", "Dans le Liptako-Gourma, elle atteint une intensité inédite (200 puis 400 tués revendiqués). L'EI-Sahel s'en sert probablement pour consolider la bande frontalière Gourma–Oudalan–Ménaka.", "Faible sur les bilans"),
    ("Initiative au nord, attrition aérienne", "Les FAMa et Africa Corps ont repris Kidal (10/10), mais perdent des moyens aériens (hélicoptère, drone MALE Orion, drone Molniya-2 abattu, drones endommagés à Sévaré) et voient monter les victimes civiles de leurs frappes. Le FLA conserve, de façon très probable, une capacité de harcèlement par drones FPV.", "Modérée"),
    ("Dommages civils : risque stratégique majeur", "Les frappes de Guèllèdjè, Tieri, Nzanakoro et Tombouctou causent 100 % des civils tués recensés. Elles nourrissent le récit du JNIM auprès des communautés peules et nomades, et donc son recrutement.", "Élevée"),
]
SCENARIOS = [
    ("Le plus probable", "55–65 %", "Usure prolongée au centre", "Le JNIM maintient 3 à 5 incidents par semaine à Ségou et Mopti (embuscades, EEI, postes isolés) et entrave les convois vers Bamako. Les FAMa et Africa Corps consolident Kidal et intensifient les frappes, avec des dommages civils récurrents.", "E2701C"),
    ("Le plus dangereux", "15–25 %", "Coup d'éclat contre un centre névralgique", "Attaque complexe contre une emprise majeure (type Dioura / Sévaré) ou la périphérie de Bamako (cercle de Kati), avec tirs indirects ou drones. Elle s'accompagne d'un durcissement du blocus des axes et d'une montée du front sud (Sikasso, frontières BF et CI).", "C8102E"),
    ("Alternatif", "≈ 20 %", "Poussée de l'EI-Sahel à l'est", "L'EI-Sahel exploite l'affaiblissement local du JNIM pour s'étendre vers Ménaka, Ansongo et Tessit et dans le Gourma. Il multiplie les embuscades contre les convois FAMa pendant que le FLA mène une guérilla par drones au nord.", "8A4FD6"),
]
INDICATORS = [
    ("Nouvelles attaques de citernes ou de convois commerciaux", "Durcissement du blocus logistique de Bamako", "Ségou–Fana, Markala, axes d'entrée de Bamako"),
    ("Incidents dans les cercles de Kati, Kolokani ou Banamba", "Encerclement rapproché de la capitale", "Koulikoro"),
    ("Tirs indirects ou drones contre aéroports et bases", "Ciblage du capacitaire aérien et ISR", "Sévaré, Gao, Tombouctou, Kidal"),
    ("Usage de drones FPV par le JNIM", "Diffusion technologique depuis le FLA", "Centre et Nord"),
    ("Attaques de postes au sud de Sikasso", "Ouverture d'un couloir vers le Golfe de Guinée", "Kadiolo, Kourouma, frontières BF / CI"),
    ("Revendications de l'EI-Sahel à Ménaka ou Ansongo", "Expansion de l'EI-Sahel vers l'ouest", "Gao, Ménaka"),
    ("Nouvelles frappes aériennes avec victimes civiles", "Ralliements communautaires et représailles", "Ségou, Mopti, Tombouctou"),
]
RECOS = [
    ("Autorités et forces de défense", [
        ("Immédiat", "Sécuriser en priorité les axes Bamako–Ségou–Niono et Ségou–Markala : ouverture d'itinéraire, lutte contre les EEI, escortes à horaires et itinéraires variables."),
        ("Immédiat", "Durcir ou regrouper les postes isolés exposés à un coup de main (scénario Kassela / Kourouma) et pré-positionner une force de réaction rapide héliportée."),
        ("Court terme", "Déployer une défense anti-drones (brouillage, détection) sur Sévaré, Goundam, Aguelhoc et Kidal ; disperser ou abriter les aéronefs."),
        ("Court terme", "Renforcer l'identification positive des cibles avant frappe, avec un mécanisme d'enquête sur les bavures : condition d'efficacité de la contre-insurrection."),
        ("Moyen terme", "Exploiter la rivalité JNIM–EI-Sahel par le renseignement (ISR) sur le Liptako, sans provoquer de rapprochement tactique entre les deux groupes."),
    ]),
    ("Partenaires, ONG et entreprises", [
        ("Immédiat", "Classer en rouge les axes Niono–Dioura, Ségou–Fana, Dougabougou–Markala, Bandiagara–Douentza et Ansongo–Ménaka : circulation de jour, en convoi, après évaluation de route la veille."),
        ("Immédiat", "Éloigner les implantations des emprises militaires et des aérodromes, cibles de tirs indirects et de drones."),
        ("Court terme", "Plans de continuité pour Bamako (carburant, eau, communications satellitaires) ; mise à jour des plans d'évacuation aérienne de Tombouctou, Gao et Kidal."),
        ("Court terme", "Éviter les points d'accès Starlink collectifs et les rassemblements en zone de frappes (la frappe de Guèllèdjè visait un point de connexion)."),
    ]),
    ("États du Golfe de Guinée", [
        ("Immédiat", "Renforcer la surveillance de la bande Sikasso – Hauts-Bassins – nord ivoirien et ghanéen : Kourouma signale une poussée du JNIM vers le sud."),
        ("Court terme", "Partage de renseignement sur les flux logistiques transfrontaliers (carburant, motos, intrants explosifs) qui alimentent les EEI."),
        ("Moyen terme", "Programmes de cohésion communautaire en zone frontalière pour couper le récit de victimisation exploité après les frappes."),
    ]),
]


def sheet_eval(wb):
    ws = wb.create_sheet("EVALUATION")
    banner(ws, "Évaluation stratégique", "Jugements établis sur l'ensemble de la base. « Probable » = 55–80 % ; « très probable » = plus de 80 %.", 8)
    widths(ws, {"A": 2, "B": 7, "C": 34, "D": 80, "E": 26, "F": 2, "G": 2, "H": 12})
    section(ws, 4, "Jugements clés", 2, 4)
    header_row(ws, 5, ["Réf.", "Jugement", "Analyse", "Niveau de confiance"], 2, TEAL)
    r = 6
    for k, (t, txt, conf) in enumerate(JUDGEMENTS, start=1):
        body(ws.cell(r, 2, f"JC{k}"), bold=True, color=TEAL, align=CENTER)
        body(ws.cell(r, 3, t), bold=True)
        body(ws.cell(r, 4, txt))
        body(ws.cell(r, 5, conf), color=MUTED, size=9)
        ws.row_dimensions[r].height = 58
        r += 1
    r += 1
    section(ws, r, "Scénarios à 30–60 jours (horizon fin décembre 2026)", 2, 4)
    r += 1
    header_row(ws, r, ["", "Scénario", "Description", "Probabilité"], 2, TEAL)
    r += 1
    for kind, prob, t, txt, col in SCENARIOS:
        body(ws.cell(r, 2, "")).fill = fill(col)
        c = body(ws.cell(r, 3, f"{kind}\n{t}"), bold=True)
        c.font = f(10, True, col)
        body(ws.cell(r, 4, txt))
        body(ws.cell(r, 5, prob), bold=True, align=CENTER, size=14, color=col)
        ws.row_dimensions[r].height = 62
        r += 1
    r += 1
    section(ws, r, "Indicateurs d'alerte à surveiller", 2, 4)
    r += 1
    header_row(ws, r, ["", "Indicateur", "Ce qu'il signalerait", "Zone"], 2, TEAL)
    r += 1
    for k, (a, b, c) in enumerate(INDICATORS, start=1):
        body(ws.cell(r, 2, f"IA{k}"), bold=True, color=TEAL, align=CENTER)
        body(ws.cell(r, 3, a), bold=True)
        body(ws.cell(r, 4, b))
        body(ws.cell(r, 5, c), color=MUTED, size=9)
        ws.row_dimensions[r].height = 32
        r += 1
    r += 1
    section(ws, r, "Recommandations", 2, 4)
    r += 1
    hz = {"Immédiat": "C8102E", "Court terme": "E2701C", "Moyen terme": "0F7C86"}
    for who, items in RECOS:
        c = ws.cell(r, 2, who)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        for col in range(2, 6):
            ws.cell(r, col).fill = fill(NAVY)
        r += 1
        for horizon, txt in items:
            body(ws.cell(r, 2, ""))
            body(ws.cell(r, 3, horizon), bold=True, color=hz[horizon])
            body(ws.cell(r, 4, txt))
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
            ws.row_dimensions[r].height = 32
            r += 1
        r += 1
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


# --------------------------------------------------------------------------- QUALITE & SOURCES
def sheet_quality(wb, incidents):
    ws = wb.create_sheet("QUALITE_SOURCES")
    banner(ws, "Qualité des données et fiabilité des sources", "Corroboration : « Corroboré » = au moins deux sources indépendantes (veille OSINT ou presse).", 9)
    widths(ws, {"A": 2, "B": 28, "C": 16, "D": 40, "E": 12, "F": 3, "G": 30, "H": 12, "I": 12})
    header_row(ws, 4, ["Niveau de fiabilité", "Incidents", "Tués associés", "Part"], 7, TEAL)
    for r, (rel, col) in enumerate(REL.items(), start=5):
        c = body(ws.cell(r, 7, rel), bold=True)
        c.font = f(10, True, col)
        body(ws.cell(r, 8, f'=COUNTIFS(inc_fiab,$G{r})'), align=CENTER)
        body(ws.cell(r, 9, f'=SUMIFS(inc_tues,inc_fiab,$G{r})'), align=CENTER)
        body(ws.cell(r, 10, f"=IF(SUM($H$5:$H$8)=0,0,H{r}/SUM($H$5:$H$8))"), fmt="0%", align=CENTER)
    ws.column_dimensions["J"].width = 10
    body(ws.cell(10, 7, "Incidents à bilan revendiqué seul"), bold=True)
    body(ws.cell(10, 8, '=COUNTIFS(inc_revend,"Oui")'), align=CENTER)
    body(ws.cell(11, 7, "Incidents géolocalisés à ± 10 km"), bold=True)
    body(ws.cell(11, 8, '=COUNTIFS(inc_prec,"Zone*")'), align=CENTER)
    pie = PieChart()
    pie.title = "Corroboration des incidents"
    pie.add_data(Reference(ws, min_col=8, min_row=4, max_row=8), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=7, min_row=5, max_row=8))
    from openpyxl.chart.series import DataPoint
    for k, col in enumerate(REL.values()):
        pt = DataPoint(idx=k)
        pt.graphicalProperties.solidFill = col
        pie.series[0].dPt.append(pt)
    pie.dataLabels = labels(val=True)
    pie.width, pie.height = 11, 7.5
    ws.add_chart(pie, "G13")
    # Sources
    header_row(ws, 4, ["Source", "Type", "Nature", "Citations"], 2)
    srcs = {}
    for i in incidents:
        for s in i["sources"]:
            srcs[s["name"]] = s
    kind_fr = {"osint": "Veille OSINT", "média": "Presse", "partisane": "Partisane", "officielle": "Officielle", "compilateur": "Compilateur"}
    kind_col = {"osint": "2F6FB5", "média": "4F8A3C", "partisane": "DC4A26", "officielle": "D69A12", "compilateur": "7D8A97"}
    rows = sorted(srcs.values(), key=lambda s: (list(kind_fr).index(s["kind"]), -sum(s["name"] in [x["name"] for x in i["sources"]] for i in incidents)))
    r = 5
    for s in rows:
        body(ws.cell(r, 2, s["name"]), bold=True)
        c = body(ws.cell(r, 3, kind_fr[s["kind"]]), bold=True)
        c.font = f(10, True, kind_col[s["kind"]])
        body(ws.cell(r, 4, s["note"]), color=MUTED, size=9)
        body(ws.cell(r, 5, f'=COUNTIFS(inc_sources,"*"&B{r}&"*")'), align=CENTER)
        r += 1
    ws.conditional_formatting.add(f"E5:E{r - 1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="0F7C86", showValue=True))
    ws.cell(r, 2, "« Citations » compte les incidents où le nom apparaît dans la colonne Sources (recherche par texte).").font = f(8, color=MUTED, italic=True)
    r += 2
    section(ws, r, "Limites et points de vigilance", 2, 4)
    r += 1
    notes = [
        ("Risque de double comptage", "INC-2026-003 (01/09, Gonadaouri / Ndaki, 200 tués) et INC-2026-025 (09/10, Ganadaouri, 400 tués) désignent des toponymes voisins. Ensemble, ils pèsent 64 % de la létalité : à recouper avant toute diffusion."),
        ("Bilans revendiqués", "Les pertes de Dioura (150) et de Sévaré (33) proviennent en partie d'AzZalaqa (JNIM) ; celles de l'EI-Sahel d'Amaq. Les lire comme des maximums."),
        ("Géolocalisation", "Plusieurs incidents sont positionnés au niveau d'une zone (une décimale, ± 10 km). INC-2026-022 (campagne de frappes sur quatre sites) est représenté par un seul point à Anéfif."),
        ("Catégories ambiguës", "INC-2026-017 : « 7 morts / blessés » russes enregistrés comme 7 morts. INC-2026-015 : attribution au JNIM présumée."),
        ("Contrôle spatial", "INC-2026-025, déclaré « Burkina Faso / Mali », se situe au Burkina Faso (Oudalan). INC-2026-003, déclaré frontalier, se situe dans le cercle de Gourma-Rharous (Tombouctou)."),
        ("Couverture", "36 incidents en 57 jours : la base reflète la sélection du compilateur, non l'exhaustivité du conflit. Interpréter avec prudence les variations de faible effectif."),
        ("Référentiel administratif", "Régions et cercles geoBoundaries (8 régions + Bamako, antérieur à la réforme de 2023). La région déclarée par la source reste disponible dans INCIDENTS."),
    ]
    for t, txt in notes:
        body(ws.cell(r, 2, t), bold=True)
        body(ws.cell(r, 3, txt))
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 44
        r += 1


# --------------------------------------------------------------------------- CARTES
def sheet_maps(wb):
    ws = wb.create_sheet("CARTES")
    banner(ws, "Cartes de situation", "Instantanés du tableau de bord interactif. Pour zoomer, filtrer et animer : liens ci-dessous.", 18)
    widths(ws, {get_column_letter(c): 9 for c in range(1, 19)})
    links = [("▶ Ouvrir la carte interactive (en ligne)", "=HYPERLINK(URL_TABLEAU,\"▶ Ouvrir la carte interactive (en ligne)\")"),
             ("▶ Ouvrir le fichier HTML local", "=HYPERLINK(FICHIER_HTML,\"▶ Ouvrir le fichier HTML local (même dossier)\")"),
             ("▶ Carte OpenStreetMap du Mali", '=HYPERLINK("https://www.openstreetmap.org/#map=6/16.5/-3.5","▶ Fond OpenStreetMap du Mali")')]
    for k, (_, fml) in enumerate(links):
        c = ws.cell(4, 2 + k * 5, fml)
        c.font = f(10, True, "0563C1", underline="single")
    cartes = ROOT / "dist" / "cartes"
    maps = [("carte_situation.png", "Carte 1 — Situation générale : incidents par acteur (taille ∝ √ tués), indice de menace régional, axes routiers, fleuve Niger, rayons de 50 / 150 / 300 km autour de Bamako."),
            ("carte_centre.png", "Carte 2 — Zoom Mali central (Koulikoro – Ségou – Mopti) : guerre des axes et pression sur Bamako."),
            ("carte_hexagones_3d.png", "Carte 3 — Densité 3D : hexagones de 30 km, hauteur et couleur = Σ(1 + √tués).")]
    r = 6
    for fn, cap in maps:
        p = cartes / fn
        c = ws.cell(r, 2, cap)
        c.font = f(10, True, NAVY)
        if p.exists():
            img = XLImage(str(p))
            img.width, img.height = 1080, int(1080 * img.height / img.width)
            ws.add_image(img, f"B{r + 1}")
            r += 1 + int(img.height / 20) + 2
        else:
            ws.cell(r + 1, 2, f"Image absente : lancez la génération des cartes ({fn}).").font = f(9, color="C8102E")
            r += 4
    src = ws.cell(r, 2, "Fonds : geoBoundaries (CC BY 4.0), Natural Earth. Données : compilations Naïm TOURÉ, 31/08 – 26/10/2026.")
    src.font = f(8, color=MUTED, italic=True)


# --------------------------------------------------------------------------- ACCUEIL
def sheet_home(wb, meta):
    ws = wb.active
    ws.title = "ACCUEIL"
    banner(ws, "OBSERVATOIRE SÉCURITAIRE · MALI — Tableau de bord décisionnel",
           f"Période du {date.fromisoformat(meta['start']):%d/%m/%Y} au {date.fromisoformat(meta['end']):%d/%m/%Y} · Sources ouvertes · Diffusion restreinte · Généré le {date.today():%d/%m/%Y}", 14)
    widths(ws, {"A": 2, **{get_column_letter(c): 13 for c in range(2, 14)}, "N": 2})
    # Indicateurs clés
    section(ws, 4, "Indicateurs clés (calculés depuis INCIDENTS)", 2, 12)
    kpis = [("Incidents", "=COUNTA(inc_id)", "13202C"), ("Tués (total)", "=SUM(inc_tues)", "13202C"), ("FDS (FAMa) tués", "=SUM(inc_fama)", "2F6FB5"),
            ("Africa Corps tués", "=SUM(inc_ac)", "2F6FB5"), ("Combattants armés tués", "=SUM(inc_ga)", "DC4A26"), ("Civils tués", "=SUM(inc_civ_k)", "C8102E"),
            ("Civils blessés", "=SUM(inc_civ_w)", "C8102E"), ("Capturés", "=SUM(inc_pow)", "13202C"), ("Véh. / aéronefs détruits", "=SUM(inc_vehd)", "13202C"),
            ("Véh. / aéronefs saisis", "=SUM(inc_vehs)", "13202C"), ("Dans le rayon de Bamako", '=COUNTIFS(inc_rayon,"Oui")', "E2701C"), ("Part du JNIM", '=IF(COUNTA(inc_id)=0,0,COUNTIFS(inc_acteur,"JNIM")/COUNTA(inc_id))', "DC4A26")]
    for k, (lab, fml, col) in enumerate(kpis):
        rr, cc = 5 + (k // 6) * 3, 2 + (k % 6) * 2
        ws.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 1)
        ws.merge_cells(start_row=rr + 1, start_column=cc, end_row=rr + 1, end_column=cc + 1)
        a = ws.cell(rr, cc, lab.upper())
        a.font = Font(name=FONT, size=8, bold=True, color=MUTED)
        a.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        v = ws.cell(rr + 1, cc, fml)
        v.font = Font(name=FONT, size=22, bold=True, color=col)
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        v.number_format = "0%" if "Part" in lab else "#,##0"
        for r2 in (rr, rr + 1):
            for c2 in (cc, cc + 1):
                ws.cell(r2, c2).fill = fill(PALE2)
                ws.cell(r2, c2).border = Border(left=Side(style="thick", color=col) if c2 == cc else None)
        ws.row_dimensions[rr + 1].height = 34
    # Accès cartographie
    section(ws, 12, "Accès à la cartographie", 2, 12)
    acc = [
        ("=HYPERLINK(URL_TABLEAU,\"▶  Ouvrir le tableau de bord et la carte interactive (en ligne)\")", "Carte WebGL, filtres, animation chronologique, dossiers d'incidents. Nécessite Internet ; lien privé à partager depuis la page."),
        ("=HYPERLINK(FICHIER_HTML,\"▶  Ouvrir la version locale (fichier HTML)\")", f"Le fichier « {HTML_FILE} » doit être dans le même dossier que ce classeur. Internet requis pour charger le moteur cartographique."),
        (None, "Voir les cartes de situation intégrées au classeur (consultables hors ligne et imprimables)."),
        (None, "Chaque incident possède des liens « OSM ↗ », « Google ↗ » et « Dossier ↗ » (colonnes AJ à AL de INCIDENTS)."),
    ]
    for k, (fml, note) in enumerate(acc):
        r = 13 + k
        c = ws.cell(r, 2)
        if fml:
            c.value = fml
        elif k == 2:
            c.value = "▶  Cartes de situation (onglet CARTES)"
            c.hyperlink = "#CARTES!A1"
        else:
            c.value = "▶  Registre des incidents avec liens cartographiques"
            c.hyperlink = "#INCIDENTS!AJ1"
        c.font = Font(name=FONT, size=11, bold=True, color="0563C1", underline="single")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        n = ws.cell(r, 8, note)
        n.font = f(9, color=MUTED)
        n.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=13)
        ws.row_dimensions[r].height = 30
    # L'essentiel
    section(ws, 18, "L'essentiel", 2, 12)
    ess = [("Le JNIM mène l'initiative.", "21 des 36 incidents (58 %), dont 16 à Ségou et Mopti ; à l'origine de 92 % des FDS maliennes tuées (192 sur 208).", "C8102E"),
           ("La pression se rapproche de Bamako.", "9 incidents à moins de 300 km, dont un poste pris à Kassela (32 km). Six attaques visent des axes routiers : asphyxie logistique.", "C8102E"),
           ("Les frappes aériennes causent toutes les morts civiles recensées.", "28 civils tués sur 28 lors de quatre frappes FAMa / Africa Corps ; ×3 entre P1 et P2 (7 → 21).", "E2701C"),
           ("La guerre JNIM – EI-Sahel saigne le Liptako.", "Deux batailles revendiquées : 600 tués, 64 % de la létalité de la base. Bilans à confirmer (double comptage possible).", "E2701C")]
    for k, (t, txt, col) in enumerate(ess):
        r = 19 + k
        a = ws.cell(r, 2, t)
        a.font = f(10, True, col)
        a.alignment = Alignment(wrap_text=True, vertical="center")
        a.border = Border(left=Side(style="thick", color=col))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        b = ws.cell(r, 6, txt)
        b.font = f(10)
        b.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
        ws.row_dimensions[r].height = 32
    # Sommaire
    section(ws, 24, "Sommaire", 2, 12)
    toc = [("CARTES", "Cartes de situation (3 cartes)"), ("INCIDENTS", "Registre enrichi des 36 incidents, liens cartographiques"), ("REGIONS", "Synthèse régionale et indice de menace"),
           ("ACTEURS", "Synthèse par acteur, matrice acteur × mode"), ("CHRONOLOGIE", "Rythme hebdomadaire, comparaison P1 / P2"), ("MODES_CIBLES", "Modes opératoires et cibles"),
           ("PRESSION_BAMAKO", "Distance à la capitale"), ("EVALUATION", "Jugements clés, scénarios, alertes, recommandations"), ("QUALITE_SOURCES", "Fiabilité, sources, limites"),
           ("PARAMETRES", "Hypothèses modifiables (période, rayon, poids de l'indice)"), ("MODE_EMPLOI", "Ouvrir la carte depuis Excel, résolution des problèmes"), ("DICTIONNAIRE", "Définition des colonnes")]
    for k, (sh, desc) in enumerate(toc):
        r, c0 = 25 + k % 6, 2 + (k // 6) * 6
        a = ws.cell(r, c0, sh)
        a.hyperlink = f"#{sh}!A1"
        a.font = f(10, True, "0563C1", underline="single")
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
        d = ws.cell(r, c0 + 2, desc)
        d.font = f(9, color=MUTED)
        ws.merge_cells(start_row=r, start_column=c0 + 2, end_row=r, end_column=c0 + 5)
    p = ROOT / "dist" / "cartes" / "carte_situation.png"
    if p.exists():
        img = XLImage(str(p))
        img.width, img.height = 700, int(700 * img.height / img.width)
        ws.add_image(img, "B32")
        ws.cell(31, 2, "Carte de situation — cliquez « Ouvrir le tableau de bord » pour la version interactive").font = f(9, True, NAVY)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --------------------------------------------------------------------------- MODE D'EMPLOI & DICTIONNAIRE
def sheet_help(wb):
    ws = wb.create_sheet("MODE_EMPLOI")
    banner(ws, "Mode d'emploi : ouvrir la carte depuis Excel", "Correctifs intégrés et résolution des problèmes courants.", 6)
    widths(ws, {"A": 2, "B": 34, "C": 100, "D": 2})
    blocks = [
        ("Ouvrir la carte", [
            ("Tableau de bord en ligne", "Onglet ACCUEIL → « Ouvrir le tableau de bord ». Le lien est privé : partagez-le depuis le menu « Partager » de la page web pour que vos destinataires puissent l'ouvrir."),
            ("Un incident précis", "INCIDENTS → colonne « Tableau de bord » (Dossier ↗) : la page s'ouvre directement sur la fiche de l'incident, centrée sur la carte."),
            ("OpenStreetMap / Google Maps", "INCIDENTS → colonnes « Carte OSM » et « Google Maps » : ouverture du point exact dans le navigateur, sans compte."),
            ("Version locale", "Placez « tableau_de_bord_mali.html » dans le même dossier que ce classeur, puis ACCUEIL → « Ouvrir la version locale ». Internet reste nécessaire pour le moteur cartographique."),
            ("Google Earth", "Double-cliquez « incidents_mali.kml » (fourni avec le classeur) : un dossier par acteur, fiche complète au clic, curseur temporel."),
            ("QGIS / uMap", "Glissez « incidents_mali.geojson » dans QGIS, ou importez-le dans uMap (umap.openstreetmap.fr). Système de coordonnées : WGS84 / EPSG:4326."),
        ]),
        ("Correctifs intégrés", [
            ("Séparateur décimal", "Dans un Excel en français, 14,5333 devient « 14,5333 » dans une URL et casse le lien. Les formules convertissent les coordonnées en texte à point décimal (SUBSTITUTE) : les liens fonctionnent quelle que soit la langue d'Excel."),
            ("Liens en formule", "Les liens sont des formules LIEN_HYPERTEXTE / HYPERLINK : si vous corrigez une latitude ou une longitude, le lien se met à jour automatiquement."),
            ("Lien direct par incident", "Le tableau de bord lit l'identifiant placé après « # » dans l'adresse (ex. …#INC-2026-007) et ouvre la fiche correspondante."),
            ("Formules extensibles", "Toutes les synthèses portent sur les lignes 2 à 500 de INCIDENTS : ajoutez de nouveaux incidents à la suite, elles se recalculent."),
            ("Compatibilité", "Fonctions standard (NB.SI.ENS / COUNTIFS, SOMME.SI.ENS / SUMIFS, INDEX / EQUIV) : Excel 2010 et suivants, Excel en ligne, LibreOffice, WPS."),
        ]),
        ("En cas de problème", [
            ("Les liens ne réagissent pas", "Le fichier s'est ouvert en « Mode protégé » (fichier reçu par courriel ou téléchargé). Cliquez « Activer la modification » dans le bandeau jaune, puis recliquez le lien."),
            ("Alerte de sécurité Microsoft", "« Les liens hypertextes peuvent être dangereux… » : cliquez « Oui ». Les liens pointent uniquement vers claude.ai, openstreetmap.org et google.com."),
            ("La version locale ne s'ouvre pas", "Vérifiez que le fichier HTML est dans le même dossier que le classeur, avec le nom exact « tableau_de_bord_mali.html ». À défaut, double-cliquez directement le fichier HTML."),
            ("Les chiffres affichent 0", "Le calcul est en mode manuel : Formules → Options de calcul → Automatique (ou touche F9)."),
            ("Fichier bloqué par Windows", "Clic droit sur le fichier → Propriétés → cocher « Débloquer » → OK (fichiers téléchargés)."),
            ("Carte vide dans le navigateur", "Le réseau bloque unpkg.com ou cdnjs.cloudflare.com (pare-feu d'entreprise). Utilisez une autre connexion ou les cartes intégrées (onglet CARTES)."),
        ]),
        ("Mettre à jour le classeur", [
            ("Nouvelle base", "Remplacez le fichier Excel source dans data/ puis lancez : python scripts/build_dashboard.py && python scripts/build_excel.py (dépôt GitHub)."),
            ("Saisie directe", "Vous pouvez aussi ajouter une ligne dans INCIDENTS : les colonnes calculées (semaine, période, total, distance, liens) se recopient en étendant les formules de la ligne précédente."),
        ]),
    ]
    r = 4
    for title, items in blocks:
        section(ws, r, title, 2, 2)
        r += 1
        for a, b in items:
            body(ws.cell(r, 2, a), bold=True)
            body(ws.cell(r, 3, b))
            ws.row_dimensions[r].height = 34
            r += 1
        r += 1

    wd = wb.create_sheet("DICTIONNAIRE")
    banner(wd, "Dictionnaire des données", "Colonnes de la feuille INCIDENTS.", 5)
    widths(wd, {"A": 2, "B": 8, "C": 28, "D": 16, "E": 90})
    header_row(wd, 4, ["Col.", "Champ", "Type", "Définition"], 2)
    defs = {
        "ID": ("Texte", "Identifiant unique de l'incident (jointure SIG, lien direct vers le tableau de bord)."),
        "Date": ("Date", "Date de l'incident."), "Semaine ISO": ("Formule", "Numéro de semaine ISO 8601."), "Période": ("Formule", "P1 ou P2 selon PARAMETRES."),
        "Pays (géocodé)": ("Texte", "Pays contenant les coordonnées (jointure spatiale)."), "Région (géocodée)": ("Texte", "Région administrative contenant les coordonnées (geoBoundaries ADM1)."),
        "Cercle (géocodé)": ("Texte", "Cercle contenant les coordonnées (geoBoundaries ADM2)."), "Région déclarée": ("Texte", "Région ou secteur indiqué par la source."),
        "Localité": ("Texte", "Lieu de l'incident."), "Latitude": ("Nombre", "WGS84, degrés décimaux."), "Longitude": ("Nombre", "WGS84, degrés décimaux."),
        "Précision géo.": ("Texte", "« Zone (±10 km) » si les coordonnées n'ont qu'une décimale, sinon « Localité »."), "Typologie (source)": ("Texte", "Typologie telle que saisie dans la base."),
        "Mode opératoire": ("Texte", "Typologie normalisée en 12 modes."), "Instigateur": ("Texte", "Famille d'acteur : JNIM, EI-Sahel, FLA, FAMa / Africa Corps, Indéterminé / accident."),
        "Instigateur (détail)": ("Texte", "Acteur tel que saisi."), "Cible": ("Texte", "Famille de cible : FAMa / Africa Corps, Civils, Milices pro-gouv., Groupes armés, Logistique civile."),
        "Cible (détail)": ("Texte", "Cible telle que saisie."), "Description": ("Texte", "Récit de l'incident."), "Morts civils": ("Nombre", ""), "Blessés civils": ("Nombre", ""),
        "Morts FDS (FAMa)": ("Nombre", "Forces de défense et de sécurité maliennes."), "Morts Africa Corps": ("Nombre", "Forces russes alliées."), "Morts groupes armés": ("Nombre", "Combattants non étatiques."),
        "Total tués": ("Formule", "Civils + FDS + Africa Corps + groupes armés."), "Capturés": ("Nombre", "Prisonniers / capturés."), "Véh./aéronefs détruits": ("Nombre", ""), "Véh./aéronefs saisis": ("Nombre", ""),
        "Armement / équipement saisi": ("Texte", ""), "Impact": ("Texte", "Impact sur les infrastructures ou le contrôle territorial."),
        "Distance Bamako (km)": ("Formule", "Distance orthodromique (haversine) au centre de Bamako."), "Dans le rayon Bamako": ("Formule", "Oui si distance ≤ rayon défini dans PARAMETRES."),
        "Fiabilité": ("Texte", "Corroboré (≥ 2 sources indépendantes), Partiellement corroboré (1), Source partisane seule, Source unique."),
        "Bilan revendiqué seul": ("Texte", "Oui si des pertes sont rapportées sans aucune source indépendante."), "Sources": ("Texte", "Sources citées, séparées par des virgules."),
        "Carte OSM": ("Lien", "Ouvre le point dans OpenStreetMap."), "Google Maps": ("Lien", "Ouvre le point dans Google Maps."), "Tableau de bord": ("Lien", "Ouvre la fiche de l'incident dans le tableau de bord interactif."),
    }
    for r, (name, _, _) in enumerate(INC_COLS, start=5):
        t, d = defs.get(name, ("", ""))
        body(wd.cell(r, 2, COL[name]), align=CENTER, color=MUTED)
        body(wd.cell(r, 3, name), bold=True)
        body(wd.cell(r, 4, t), color="0F7C86" if t == "Formule" else INK)
        body(wd.cell(r, 5, d))


def netlify_pack():
    """Dossier prêt à glisser sur Netlify Drop : index.html + fichiers téléchargeables, non indexé."""
    import shutil
    import zipfile

    dist = ROOT / "dist"
    site = dist / "site_netlify"
    shutil.rmtree(site, ignore_errors=True)
    (site / "fichiers").mkdir(parents=True)
    shutil.copy(dist / HTML_FILE, site / "index.html")
    for fn in ("Tableau_de_bord_Mali.xlsx", "incidents_mali.kml", "incidents_mali.geojson", "incidents_mali_enrichi.csv"):
        shutil.copy(dist / fn, site / "fichiers" / fn)
    # Données sensibles : on demande aux moteurs de recherche de ne pas indexer le site.
    (site / "_headers").write_text("/*\n  X-Robots-Tag: noindex, nofollow\n  Referrer-Policy: no-referrer\n", encoding="utf-8")
    (site / "robots.txt").write_text("User-agent: *\nDisallow: /\n", encoding="utf-8")
    with zipfile.ZipFile(dist / "Site_Netlify_Mali.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in site.rglob("*"):
            if f.is_file():
                z.write(f, f.relative_to(site))
    print(dist / "Site_Netlify_Mali.zip")


def main(xlsx):
    geo, incidents, d0, d1, split = build(xlsx)
    meta = {"start": d0.isoformat(), "end": d1.isoformat(), "split": split.isoformat(), "bamako": {"lat": 12.6392, "lon": -8.0029}}
    wb = Workbook()
    sheet_home(wb, meta)
    sheet_maps(wb)
    sheet_incidents(wb, incidents)
    inc_names(wb)
    sheet_regions(wb, incidents)
    sheet_actors(wb, incidents)
    sheet_chrono(wb, meta)
    sheet_modes(wb, incidents)
    sheet_bamako(wb, incidents)
    sheet_eval(wb)
    sheet_quality(wb, incidents)
    sheet_params(wb, meta)
    sheet_help(wb)
    colors = {"ACCUEIL": NAVY, "CARTES": TEAL, "EVALUATION": "C8102E", "PARAMETRES": "0000FF", "MODE_EMPLOI": "7D8A97", "DICTIONNAIRE": "7D8A97"}
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = colors.get(ws.title, ws.sheet_properties.tabColor or "566577")
        ws.sheet_view.zoomScale = 90
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.left.text = "Observatoire sécuritaire Mali · Sources ouvertes · Diffusion restreinte"
        ws.oddFooter.right.text = "&A · page &P / &N"
    wb.properties.title = "Observatoire sécuritaire Mali — tableau de bord décisionnel"
    wb.properties.creator = "Observatoire sécuritaire Mali"
    wb.calculation.fullCalcOnLoad = True  # Excel recalcule toutes les formules à l'ouverture
    wb.save(OUT)
    print(OUT)
    netlify_pack()


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX)
